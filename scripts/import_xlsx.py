#!/usr/bin/env python3
"""Generate supabase/seed.sql from "Tracker Tindak Lanjut Uplifting JOUMPA.xlsx".

    python3 -I scripts/import_xlsx.py [--xlsx PATH] [--seed PATH] [--rows PATH] [--check]

Writes
  supabase/seed.sql      staff (+ bmi_note), action_items (+ due_rule), replacements
                         (+ staff_code). Idempotent:
                         staff/action_items upsert on their code; replacements MERGE on
                         (report_group, replaced_name). One transaction. Does not touch
                         stations/settings/sync_state (migration) or scores/BMI checks.
  scripts/out/rows.json  the same rows in camelCase + settings + the workbook's cached
                         per-row results, for the parity test (lib/workbook.test.ts).
Both contain personal data and are gitignored (scripts/out/ carries its own .gitignore).

Reads cached values (data_only=True) for inputs, so the few formula inputs (rolling due
dates of TL10/TL12/TL13) are stored as the workbook last computed them. Output is
deterministic: same workbook bytes -> same files.

--check only validates and prints the data report, writing nothing.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import unicodedata
import sys
import warnings
from collections import Counter, defaultdict
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
import openpyxl  # noqa: E402
from openpyxl.utils.datetime import from_excel  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "Tracker Tindak Lanjut Uplifting JOUMPA.xlsx"
DEFAULT_SEED = ROOT / "supabase" / "seed.sql"
DEFAULT_ROWS = ROOT / "scripts" / "out" / "rows.json"

STATIONS = ("SUB", "DPS", "CGK", "HLP", "KNO")
REPORT_GROUPS = ("SUB", "DPS", "CGK & HLP", "KNO")
CONCLUSIONS = ("Sesuai", "Sesuai dengan Catatan", "Perlu Perbaikan", "Tidak Sesuai")
ASSIGNMENT = ("Aktif", "Coaching 30 Hari", "Diganti", "Ditarik")
ACTION_STATUS = ("Belum Mulai", "On Progress", "Selesai", "Tertunda")
KINDS = ("Sekali", "Rutin")
GENDERS = ("L", "P")
REPORTED = ("Ya", "Belum")

SCORE_COLS = ("score_a", "score_b", "score_c", "score_d", "score_e", "score_f")
BMI_PERIOD_FIRST_COL = 6  # F: Tgl Cek of period 1; 5 columns per period (date, height, weight, BMI, category)


class ImportError_(Exception):
    pass


# --------------------------------------------------------------------------- value cleaning


def fail(where: str, msg: str) -> None:
    raise ImportError_(f"{where}: {msg}")


def text(v, where: str) -> str | None:
    """Empty or whitespace-only -> None; strings are trimmed; integral numbers -> '123'."""
    if v is None:
        return None
    if isinstance(v, bool):
        fail(where, f"unexpected boolean {v!r}")
    if isinstance(v, (int, float)):
        if isinstance(v, float) and not v.is_integer():
            return repr(v)
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        fail(where, f"unexpected date {v!r} in a text column")
    s = str(v).strip()
    return s or None


def choice(v, allowed: tuple[str, ...], where: str) -> str | None:
    s = text(v, where)
    if s is None:
        return None
    for a in allowed:
        if s.upper() == a.upper():
            return a
    fail(where, f"{s!r} is not one of {allowed}")


def number(v, where: str):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        fail(where, f"unexpected boolean {v!r}")
    if isinstance(v, str):
        try:
            v = float(v.strip().replace(",", "."))
        except ValueError:
            fail(where, f"not a number: {v!r}")
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v.is_integer():
            return int(v)
        return v
    fail(where, f"not a number: {v!r}")


def score(v, where: str) -> int | None:
    n = number(v, where)
    if n is None:
        return None
    if n != int(n) or not 1 <= n <= 5:
        fail(where, f"score must be a whole number 1-5, got {n!r}")
    return int(n)


def iso_date(v, where: str) -> str | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, dt.datetime):
        if v.time() != dt.time(0, 0):
            fail(where, f"date has a time part: {v!r}")
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return from_excel(v).date().isoformat()
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
            try:
                return dt.datetime.strptime(s, fmt).date().isoformat()
            except ValueError:
                pass
    fail(where, f"not a date: {v!r}")


def progress_pct(v, where: str) -> int:
    """Excel % Progres is a fraction 0..1; the DB stores an integer 0..100."""
    n = number(v, where)
    if n is None:
        return 0
    if not 0 <= n <= 1:
        fail(where, f"progress must be a fraction 0..1 (Excel %), got {n!r}")
    return int((Decimal(repr(n)) * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def camel(key: str) -> str:
    return re.sub(r"_([a-z0-9])", lambda m: m.group(1).upper(), key)


# --------------------------------------------------------------------------- SQL


def sql_literal(v, pgtype: str | None = None) -> str:
    cast = f"::{pgtype}" if pgtype else ""
    if v is None:
        return "null" + cast
    if isinstance(v, bool):
        return ("true" if v else "false") + cast
    if isinstance(v, (int, float)):
        return (str(v) if isinstance(v, int) else repr(v)) + cast
    s = str(v)
    if "\x00" in s:
        raise ImportError_("NUL character in text value")
    return "'" + s.replace("'", "''") + "'" + cast


def upsert(table: str, key: str, cols: list[str], rows: list[dict]) -> str:
    """INSERT ... VALUES coerces literals to the column types, so no casts are needed."""
    lines = [f"insert into public.{table} ({', '.join(cols)}) values"]
    body = []
    for r in rows:
        body.append("  (" + ", ".join(sql_literal(r[c]) for c in cols) + ")")
    lines.append(",\n".join(body))
    updates = ",\n  ".join(f"{c} = excluded.{c}" for c in cols if c != key)
    lines.append(f"on conflict ({key}) do update set\n  {updates};")
    return "\n".join(lines)


def merge(table: str, match: list[str], cols: list[str], rows: list[dict], types: dict[str, str]) -> str:
    """MERGE keeps identity ids stable on re-runs and needs no unique constraint (PG 15+)."""
    values = ",\n    ".join(
        "(" + ", ".join(sql_literal(r[c], types[c]) for c in cols) + ")" for r in rows
    )
    on = " and ".join(f"t.{c} is not distinct from s.{c}" for c in match)
    sets = ",\n    ".join(f"{c} = s.{c}" for c in cols if c not in match)
    return (
        f"merge into public.{table} as t\n"
        f"using (values\n    {values}\n) as s({', '.join(cols)})\n"
        f"on {on}\n"
        f"when matched then update set\n    {sets}\n"
        f"when not matched then insert ({', '.join(cols)})\n"
        f"  values ({', '.join('s.' + c for c in cols)});"
    )


# --------------------------------------------------------------------------- workbook


def cell(ws, row: int, col: int):
    return ws.cell(row=row, column=col).value


def read_settings(wv, wf) -> dict:
    p = wv["Parameter"]
    where = "Parameter"
    deadline_formula = str(wf["Penggantian SDM"]["M5"].value or "")
    m = re.search(r"DATE\((\d{4}),(\d{1,2}),(\d{1,2})\)", deadline_formula)
    if not m:
        fail("Penggantian SDM!M5", "replacement deadline DATE(y,m,d) not found in formula")
    deadline = dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
    return {
        "week1Start": iso_date(p["C5"].value, f"{where}!C5"),
        "weeks": number(p["C6"].value, f"{where}!C6"),
        "bmiFirstCheck": iso_date(p["C7"].value, f"{where}!C7"),
        "bmiIntervalDays": number(p["C8"].value, f"{where}!C8"),
        "bmiPeriods": number(p["C9"].value, f"{where}!C9"),
        "bmiUnderweightBelow": number(p["C11"].value, f"{where}!C11"),
        "bmiNormalMax": number(p["C12"].value, f"{where}!C12"),
        "bmiOverweightMax": number(p["C13"].value, f"{where}!C13"),
        "minHeightFemale": number(p["C16"].value, f"{where}!C16"),
        "minHeightMale": number(p["C17"].value, f"{where}!C17"),
        "passAvgMin": number(p["C19"].value, f"{where}!C19"),
        "improveAvgMin": number(p["C20"].value, f"{where}!C20"),
        "posttestMin": number(p["C21"].value, f"{where}!C21"),
        "replacementDeadline": deadline,
    }


def read_staff(wv) -> tuple[list[dict], dict[str, dict], list[str]]:
    ws = wv["Master SDM"]
    bmi = wv["Cek BMI 2 Mingguan"]
    rows, cached, notes = [], {}, []
    for r in range(5, ws.max_row + 1):
        w = f"Master SDM!{r}"
        code = text(cell(ws, r, 2), f"{w} B")
        name = text(cell(ws, r, 5), f"{w} E")
        if code is None and name is None:
            continue
        if name is None:
            # Empty TMB-xx placeholder rows. Anything else on such a row would be lost.
            stray = [c for c in (3, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 17, 20, 21) if cell(ws, r, c) not in (None, "")]
            stray = [c for c in stray if not (c == 20 and cell(ws, r, c) == "Aktif")]
            if stray:
                fail(w, f"row {code} has no name but has data in columns {stray}")
            continue
        if code is None:
            fail(w, "row has a name but no ID SDM")
        # Cek BMI rows mirror Master SDM rows one-to-one (BMI row = Master row + 1).
        bmi_code = cell(bmi, r + 1, 2)
        if bmi_code != code:
            fail(f"Cek BMI 2 Mingguan!B{r + 1}", f"expected {code}, found {bmi_code!r}")
        station = choice(cell(ws, r, 3), STATIONS, f"{w} C")
        row = {
            "code": code,
            "station": station,
            "name": name,
            "nipp": text(cell(ws, r, 6), f"{w} F"),
            "gender": choice(cell(ws, r, 7), GENDERS, f"{w} G"),
            "pre_test": number(cell(ws, r, 8), f"{w} H"),
            "post_test": number(cell(ws, r, 9), f"{w} I"),
            **{c: score(cell(ws, r, 10 + i), f"{w} {'JKLMNO'[i]}") for i, c in enumerate(SCORE_COLS)},
            "report_conclusion": choice(cell(ws, r, 17), CONCLUSIONS, f"{w} Q"),
            "assignment_status": choice(cell(ws, r, 20), ASSIGNMENT, f"{w} T") or "Aktif",
            "notes": text(cell(ws, r, 21), f"{w} U"),
            "bmi_note": text(cell(bmi, r + 1, 37), f"Cek BMI 2 Mingguan!AK{r + 1}"),
            "sort_order": len(rows) + 1,
        }
        rows.append(row)
        cached[code] = {
            "practiceAvg": cell(ws, r, 16) if cell(ws, r, 16) != "" else None,
            "criteriaStatus": cell(ws, r, 18) or None,
            "consistency": cell(ws, r, 19) or None,
        }
        # BMI inputs: 5 periods x (date, height, weight). The workbook has none yet.
        for p in range(5):
            c0 = BMI_PERIOD_FIRST_COL + 5 * p
            if any(cell(bmi, r + 1, c0 + k) not in (None, "") for k in range(3)):
                notes.append(f"BMI period {p + 1} has data for {code}")
    return rows, cached, notes


def read_weekly_scores(wv) -> list[str]:
    ws = wv["Log Performa Mingguan"]
    found = []
    for r in range(5, ws.max_row + 1):
        if any(cell(ws, r, c) not in (None, "") for c in list(range(6, 12)) + [15, 16]):
            found.append(f"Log Performa Mingguan row {r} has scores/notes")
    return found


DUE_RULES = ("bulanan-tgl-5", "cek-bmi-berikutnya", "jumat-berikutnya")


def infer_due_rule(schedule: str | None) -> str | None:
    """Same logic as inferDueRule() in lib/rules.ts: the Jadwal text states the rolling rule."""
    if not schedule:
        return None
    t = schedule.lower()
    if not re.search(r"batas\s*=", t):
        return None
    if re.search(r"batas\s*=\s*jatuh tempo berikutnya", t) and re.search(r"tanggal\s+5\b", t):
        return "bulanan-tgl-5"
    if re.search(r"batas\s*=\s*cek berikutnya", t):
        return "cek-bmi-berikutnya"
    if re.search(r"batas\s*=\s*jumat berikutnya", t):
        return "jumat-berikutnya"
    return None


def formula_due_rule(formula) -> str | None:
    """Which rolling rule a Batas Waktu formula implements (None for a typed date)."""
    if not (isinstance(formula, str) and formula.startswith("=")):
        return None
    f = formula.upper().replace("$", "")
    if "DAY(PARAMETER!C23)>5" in f:
        return "bulanan-tgl-5"
    if "ROUNDUP(" in f and "PARAMETER!C8" in f:
        return "cek-bmi-berikutnya"
    if "WEEKDAY(" in f:
        return "jumat-berikutnya"
    return "unknown formula"


def read_action_items(wv, wf) -> tuple[list[dict], dict[str, dict]]:
    ws, wsf = wv["Tindak Lanjut"], wf["Tindak Lanjut"]
    rows, cached = [], {}
    for r in range(5, ws.max_row + 1):
        w = f"Tindak Lanjut!{r}"
        code = text(cell(ws, r, 2), f"{w} B")
        if code is None:
            if any(cell(ws, r, c) not in (None, "") for c in range(3, 16)):
                fail(w, "row has data but no ID")
            continue
        due_formula = cell(wsf, r, 9)
        row = {
            "code": code,
            "report_group": choice(cell(ws, r, 3), REPORT_GROUPS, f"{w} C"),
            "area": text(cell(ws, r, 4), f"{w} D"),
            "action": text(cell(ws, r, 5), f"{w} E"),
            "target": text(cell(ws, r, 6), f"{w} F"),
            "kind": choice(cell(ws, r, 7), KINDS, f"{w} G"),
            "schedule": text(cell(ws, r, 8), f"{w} H"),
            "due_date": iso_date(cell(ws, r, 9), f"{w} I"),
            "pic": text(cell(ws, r, 10), f"{w} J"),
            "status": choice(cell(ws, r, 11), ACTION_STATUS, f"{w} K") or "Belum Mulai",
            "progress": progress_pct(cell(ws, r, 12), f"{w} L"),
            "updated_on": iso_date(cell(ws, r, 13), f"{w} M"),
            "evidence": text(cell(ws, r, 14), f"{w} N"),
            "kps_notes": text(cell(ws, r, 15), f"{w} O"),
            "due_rule": infer_due_rule(text(cell(ws, r, 8), f"{w} H")),
            "sort_order": len(rows) + 1,
        }
        # The schedule text and the Batas Waktu formula must agree on the rolling rule.
        from_formula = formula_due_rule(due_formula)
        if from_formula != row["due_rule"]:
            fail(f"{w} H/I", f"{code}: schedule says due_rule={row['due_rule']!r}, Batas Waktu formula says {from_formula!r}")
        rows.append(row)
        cached[code] = {
            "dueDate": row["due_date"],
            "dueIsFormula": isinstance(due_formula, str) and due_formula.startswith("="),
            "daysLeft": cell(ws, r, 16) if cell(ws, r, 16) != "" else None,
            "flag": cell(ws, r, 17) or None,
        }
    return rows, cached


def norm_name(s: str) -> str:
    s = unicodedata.normalize("NFKC", s).casefold()
    s = re.sub(r"[.,'’`]", " ", s)
    return " ".join(s.split())


def match_staff(row: dict, staff: list[dict]) -> tuple[str | None, str]:
    """staff_code for a replacement row: unique normalized-name match within its station
    (or report group). Tries the full name, then the part before an annotation ('/' or '(').
    Returns (code or None, how)."""
    def group_of(station: str) -> str:
        return "CGK & HLP" if station in ("CGK", "HLP") else station

    if row["station"]:
        pool = [s for s in staff if s["station"] == row["station"]]
    else:
        pool = [s for s in staff if group_of(s["station"]) == row["report_group"]]
    attempts = [("exact", row["replaced_name"])]
    head = re.split(r"[/(]", row["replaced_name"], maxsplit=1)[0]
    if head != row["replaced_name"]:
        attempts.append(("before '/' or '('", head))
    for how, candidate in attempts:
        key = norm_name(candidate)
        if not key:
            continue
        hits = [s["code"] for s in pool if norm_name(s["name"]) == key]
        if len(hits) == 1:
            return hits[0], how
        if len(hits) > 1:
            return None, f"ambiguous ({how}): {hits}"
    elsewhere = [s["code"] for s in staff if s not in pool and norm_name(s["name"]) in {norm_name(c) for _, c in attempts}]
    return None, f"no match in {row['station'] or row['report_group']}" + (f" (same name elsewhere: {elsewhere})" if elsewhere else "")


def read_replacements(wv) -> tuple[list[dict], list[dict]]:
    ws = wv["Penggantian SDM"]
    rows, cached = [], []
    for r in range(5, ws.max_row + 1):
        w = f"Penggantian SDM!{r}"
        values = [cell(ws, r, c) for c in range(2, 15) if c not in (12, 13)]
        if all(v in (None, "") for v in values):
            continue
        replaced = text(cell(ws, r, 4), f"{w} D")
        if replaced is None:
            fail(w, "row has data but no 'Nama SDM Diganti' (replaced_name is NOT NULL)")
        row = {
            "report_group": choice(cell(ws, r, 2), REPORT_GROUPS, f"{w} B"),
            "station": choice(cell(ws, r, 3), STATIONS, f"{w} C"),
            "replaced_name": replaced,
            "reason": text(cell(ws, r, 5), f"{w} E"),
            "withdrawn_on": iso_date(cell(ws, r, 6), f"{w} F"),
            "replacement_name": text(cell(ws, r, 7), f"{w} G"),
            "effective_on": iso_date(cell(ws, r, 8), f"{w} H"),
            "training_on": iso_date(cell(ws, r, 9), f"{w} I"),
            "post_test": number(cell(ws, r, 10), f"{w} J"),
            "practice_avg": number(cell(ws, r, 11), f"{w} K"),
            "reported": choice(cell(ws, r, 14), REPORTED, f"{w} N"),
            "notes": text(cell(ws, r, 15), f"{w} O"),
            "staff_code": None,  # set by link_replacements()
            "sort_order": len(rows) + 1,
        }
        rows.append(row)
        cached.append({"pass": cell(ws, r, 12) or None, "timeliness": cell(ws, r, 13) or None})
    return rows, cached


def link_replacements(repl: list[dict], staff: list[dict]) -> list[str]:
    """Fills staff_code; replaced_name keeps the original text. Returns log lines (IDs only)."""
    lines = []
    for i, r in enumerate(repl, start=1):
        code, how = match_staff(r, staff)
        r["staff_code"] = code
        lines.append(f"replacement #{i} -> {code} ({how})" if code else f"WARNING replacement #{i}: staff_code left NULL, {how}")
    return lines


# --------------------------------------------------------------------------- report


def data_report(staff: list[dict], cached_staff: dict, items: list[dict], repl: list[dict]) -> list[str]:
    """Things that look wrong in the data. IDs only, never names."""
    out = []
    by_station = Counter(s["station"] for s in staff)
    out.append("staff per station: " + ", ".join(f"{k} {by_station.get(k, 0)}" for k in STATIONS))
    nipp = defaultdict(list)
    for s in staff:
        if s["nipp"]:
            nipp[s["nipp"]].append(s["code"])
    dup = {k: v for k, v in nipp.items() if len(v) > 1}
    out.append(f"duplicate NIPP: {dup or 'none'}")
    out.append(f"NIPP missing: {sum(1 for s in staff if not s['nipp'])} of {len(staff)}")
    out.append(f"gender (L/P) missing: {sum(1 for s in staff if not s['gender'])} of {len(staff)}")
    no_post = [s["code"] for s in staff if s["post_test"] is None]
    out.append(f"no post-test: {len(no_post)} {no_post}")
    low_post = [s["code"] for s in staff if s["post_test"] is not None and s["post_test"] < 80]
    out.append(f"post-test < 80: {low_post}")
    drop = [s["code"] for s in staff if s["pre_test"] is not None and s["post_test"] is not None and s["post_test"] < s["pre_test"]]
    out.append(f"post-test lower than pre-test: {drop}")
    beda = [c for c, v in cached_staff.items() if str(v["consistency"] or "").startswith("Beda")]
    out.append(f"report conclusion != criteria 6.2 ('Beda'): {len(beda)} {beda}")
    names = {s["name"] for s in staff}
    unmatched = [i + 1 for i, r in enumerate(repl) if r["replaced_name"] not in names]
    out.append(f"replacements whose replaced_name is not a roster name verbatim (row #): {unmatched}")
    out.append(f"action items with a rolling due_rule: {sum(1 for it in items if it['due_rule'])} "
               + str(dict(Counter(it["due_rule"] for it in items if it["due_rule"]))))
    out.append(f"replacements with reported empty (row #): {[i + 1 for i, r in enumerate(repl) if r['reported'] is None]}")
    runs: list[list] = []
    for s in staff:
        if runs and runs[-1][0] == s["station"]:
            runs[-1][1] += 1
        else:
            runs.append([s["station"], 1])
    if len(runs) > len({r[0] for r in runs}):
        out.append("station blocks in sheet order (interleaved): " + ", ".join(f"{k}x{n}" for k, n in runs))
    return out


# --------------------------------------------------------------------------- main


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--seed", type=Path, default=DEFAULT_SEED)
    ap.add_argument("--rows", type=Path, default=DEFAULT_ROWS)
    ap.add_argument("--check", action="store_true", help="validate and report only")
    args = ap.parse_args(argv)

    raw = args.xlsx.read_bytes()
    sha = hashlib.sha256(raw).hexdigest()
    wv = openpyxl.load_workbook(args.xlsx, data_only=True)
    wf = openpyxl.load_workbook(args.xlsx, data_only=False)

    settings = read_settings(wv, wf)
    staff, cached_staff, bmi_data = read_staff(wv)
    score_data = read_weekly_scores(wv)
    if bmi_data or score_data:
        fail("workbook", "weekly scores / BMI checks now contain data; extend the generator to import them: "
             + "; ".join((bmi_data + score_data)[:5]))
    items, cached_items = read_action_items(wv, wf)
    repl, cached_repl = read_replacements(wv)
    links = link_replacements(repl, staff)
    today = iso_date(wv["Parameter"]["C23"].value, "Parameter!C23")

    report = data_report(staff, cached_staff, items, repl) + links
    print(f"workbook sha256 {sha[:16]}…  cached TODAY() = {today}", file=sys.stderr)
    print(f"staff {len(staff)}, action_items {len(items)}, replacements {len(repl)}, "
          f"bmi_note {sum(1 for s in staff if s['bmi_note'])}", file=sys.stderr)
    for line in report:
        print("  - " + line, file=sys.stderr)
    if args.check:
        return 0

    staff_cols = ["code", "station", "name", "nipp", "gender", "pre_test", "post_test", *SCORE_COLS,
                  "report_conclusion", "assignment_status", "notes", "bmi_note", "sort_order"]
    item_cols = ["code", "report_group", "area", "action", "target", "kind", "schedule", "due_date", "pic",
                 "status", "progress", "updated_on", "evidence", "kps_notes", "due_rule", "sort_order"]
    repl_cols = ["report_group", "station", "replaced_name", "reason", "withdrawn_on", "replacement_name",
                 "effective_on", "training_on", "post_test", "practice_avg", "reported", "notes", "staff_code",
                 "sort_order"]
    types = {
        **{c: "text" for c in staff_cols + item_cols + repl_cols},
        **{c: "numeric" for c in ("pre_test", "post_test", "practice_avg")},
        **{c: "smallint" for c in SCORE_COLS + ("progress",)},
        **{c: "date" for c in ("due_date", "updated_on", "withdrawn_on", "effective_on", "training_on")},
        "sort_order": "integer",
    }
    sql = "\n".join([
        "-- Generated by scripts/import_xlsx.py from \"Tracker Tindak Lanjut Uplifting JOUMPA.xlsx\"",
        f"-- workbook sha256 {sha}",
        "-- Contains personal data (names, NIPP). Gitignored; do not commit. Do not edit by hand.",
        "-- Idempotent: re-running resets these rows to the workbook values; other rows are untouched.",
        "",
        "set client_encoding = 'UTF8';",
        "begin;",
        "",
        f"-- staff: {len(staff)} rows (Master SDM; bmi_note from Cek BMI 2 Mingguan AK)",
        upsert("staff", "code", staff_cols, staff),
        "",
        f"-- action_items: {len(items)} rows (Tindak Lanjut; progress = Excel fraction x 100;"
        f" due_rule on {sum(1 for it in items if it['due_rule'])} rolling items)",
        upsert("action_items", "code", item_cols, items),
        "",
        f"-- replacements: {len(repl)} rows (Penggantian SDM; staff_code linked by normalized name)",
        merge("replacements", ["report_group", "replaced_name"], repl_cols, repl, types),
        "",
        "commit;",
        "",
    ])
    args.seed.parent.mkdir(parents=True, exist_ok=True)
    args.seed.write_text(sql, encoding="utf-8")

    def camel_rows(rows):
        return [{camel(k): v for k, v in r.items()} for r in rows]

    out_dir = args.rows.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    gi = out_dir / ".gitignore"
    if out_dir.resolve() == DEFAULT_ROWS.parent.resolve() and not gi.exists():
        gi.write_text("# generated from the workbook: personal data, never commit\n*\n", encoding="utf-8")
    payload = {
        "source": args.xlsx.name,
        "sha256": sha,
        "today": today,
        "settings": settings,
        "staff": camel_rows(staff),
        "weeklyScores": [],
        "bmiChecks": [],
        "actionItems": camel_rows(items),
        "replacements": camel_rows(repl),
        "cached": {"staff": cached_staff, "actionItems": cached_items, "replacements": cached_repl},
    }
    args.rows.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {args.seed.relative_to(ROOT) if args.seed.is_relative_to(ROOT) else args.seed} "
          f"and {args.rows.relative_to(ROOT) if args.rows.is_relative_to(ROOT) else args.rows}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main(sys.argv[1:]))
    except ImportError_ as e:
        print(f"import_xlsx: {e}", file=sys.stderr)
        sys.exit(1)
